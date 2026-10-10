"""Check the built Windows executable, including its bundled dependencies."""
import subprocess
import tempfile
import time
from pathlib import Path

from docx import Document


exe = Path("dist/WordToTxt.exe").resolve()
with tempfile.TemporaryDirectory() as directory:
    src = Path(directory) / "Тест с пробелами.docx"
    doc = Document()
    doc.add_paragraph("Проверочная тема")
    doc.add_paragraph("1. Сколько будет два плюс два?")
    doc.add_paragraph("А) Три")
    doc.add_paragraph("Б) Четыре").runs[0].bold = True
    doc.save(src)
    process = subprocess.Popen([str(exe), str(src)])
    try:
        result = src.with_suffix(".txt")
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            if result.exists():
                text = result.read_text(encoding="utf-8")
                if "CORRECT: B\n" in text:
                    break
            if process.poll() is not None:
                raise RuntimeError(f"EXE exited before producing a result: {process.returncode}")
            time.sleep(0.2)
        else:
            raise RuntimeError("EXE did not produce the expected TXT within 60 seconds")
        for expected in (
            "TEXT: Сколько будет два плюс два?",
            "TOPIC: Проверочная тема",
            "A: Три",
            "B: Четыре",
        ):
            if expected not in text:
                raise RuntimeError(f"Missing output: {expected}")
        print("Windows EXE successfully converted DOCX to UTF-8 TXT")
    finally:
        # Close the report dialog and the onefile bootloader's child process.
        if process.poll() is None:
            subprocess.run(
                ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                check=False, capture_output=True,
            )
        process.wait(timeout=15)
